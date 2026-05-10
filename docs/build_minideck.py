"""
Build a compact 3-slide standalone deck containing only the three quantum
use-case slides from the full executive deck:

    1. Six use cases at HDFC + an honest timeline
    2. Use case #1 - Credit-card fraud detection (deep dive)
    3. Use case #2 - Anti-Money Laundering / AML (deep dive)

It re-runs `build_deck.py` to regenerate the full 22-slide deck, then strips
out all but the three target slides and renumbers the footer slide indices
(11/12/13 -> 01/02/03).

Usage:
    pip install python-pptx
    python docs/build_minideck.py
    # writes docs/HDFC_Quantum_UseCases_MiniDeck.pptx

Set REGEN=0 in the environment to skip regenerating the full deck (use the
existing one on disk).
"""
import os
import subprocess
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

DOCS = os.path.dirname(os.path.abspath(__file__))
FULL = os.path.join(DOCS, "HDFC_Quantum_Executive_Deck.pptx")
MINI = os.path.join(DOCS, "HDFC_Quantum_UseCases_MiniDeck.pptx")
BUILD = os.path.join(DOCS, "build_deck.py")

# 0-based indices of the three slides to keep (1-based: slides 11, 12, 13).
KEEP_INDICES = [10, 11, 12]

# Footer text replacements: old slide number (str) -> new slide number (str).
# Use zero-padded form to match the formatting in build_deck.py footer().
FOOTER_RENUMBER = {"11": "01", "12": "02", "13": "03"}


def _regen_full_deck() -> None:
    """Re-run the main deck generator unless REGEN=0 is set."""
    if os.environ.get("REGEN", "1") != "0":
        subprocess.check_call([sys.executable, BUILD])


def _drop_slides(prs: Presentation, keep: list[int]) -> None:
    """Remove every slide whose index is not in `keep` and drop its
    relationship from the package so the resulting .pptx is self-consistent.
    """
    sld_id_lst = prs.slides._sldIdLst
    slide_elements = list(sld_id_lst)

    keep_set = set(keep)
    for idx, sld_id in enumerate(slide_elements):
        if idx in keep_set:
            continue
        r_id = sld_id.get(qn("r:id"))
        sld_id_lst.remove(sld_id)
        if r_id is not None:
            prs.part.drop_rel(r_id)


def _renumber_footers(prs: Presentation, mapping: dict[str, str]) -> None:
    """Walk each surviving slide's text shapes and replace footer slide
    numbers (e.g. "11") with their mini-deck equivalents (e.g. "01").

    The footer in build_deck.py uses zero-padded two-digit numbers placed in
    a small right-aligned text box at the bottom-right of the slide; this
    matcher only rewrites exact-match runs to avoid touching body content.
    """
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            for paragraph in tf.paragraphs:
                for run in paragraph.runs:
                    text = run.text.strip()
                    if text in mapping:
                        run.text = mapping[text]


def main() -> None:
    _regen_full_deck()

    if not os.path.exists(FULL):
        raise SystemExit(
            f"Expected {FULL!r} to exist after running build_deck.py."
        )

    prs = Presentation(FULL)
    _drop_slides(prs, KEEP_INDICES)
    _renumber_footers(prs, FOOTER_RENUMBER)
    prs.save(MINI)
    print(f"WROTE {MINI}")


if __name__ == "__main__":
    main()
